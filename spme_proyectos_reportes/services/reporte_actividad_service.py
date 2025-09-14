# services/reporte_actividad_service.py
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.conf import settings
from django.utils import timezone
from datetime import timedelta
import tempfile
import os

class ReporteActividadService:
    def __init__(self, actividad):
        self.actividad = actividad
        self.document = Document()
        self.setup_styles()
    
    def setup_styles(self):
        """Configura los estilos del documento"""
        styles = self.document.styles
        style = styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(11)
    
    def generar_reporte_actividad(self):
        """Genera el reporte completo de la actividad"""
        try:
            self.agregar_portada()
            self.agregar_informacion_basica()
            self.agregar_informacion_financiera()
            self.agregar_fechas_estado()
            self.agregar_tareas()
            self.agregar_relaciones()
            
            return self.guardar_documento_temp()
            
        except Exception as e:
            raise Exception(f"Error generando reporte: {str(e)}")
    
    def agregar_portada(self):
        """Agrega la portada del reporte"""
        title = self.document.add_heading('REPORTE DE ACTIVIDAD', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        for run in title.runs:
            run.font.color.rgb = RGBColor(0, 0, 139)
        
        self.document.add_paragraph()
        self.document.add_paragraph(f'Código: {self.actividad.codigo or "N/A"}')
        self.document.add_paragraph(f'Nombre: {self.actividad.nombreCorto or "N/A"}')
        
        if self.actividad.proyecto:
            self.document.add_paragraph(f'Proyecto: {self.actividad.proyecto.codigo} - {self.actividad.proyecto.titulo}')
        
        gmt_minus_4 = timezone.now() - timedelta(hours=4)
        self.document.add_paragraph(f'Fecha de generación: {gmt_minus_4.strftime("%Y-%m-%d %H:%M:%S")} (GMT-4)')
        
        self.document.add_page_break()
    
    def agregar_informacion_basica(self):
        """Agrega información básica de la actividad"""
        self.document.add_heading('INFORMACIÓN BÁSICA', level=1)
        
        table = self.document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        datos_basicos = [
            ('Código', self.actividad.codigo),
            ('Nombre Corto', self.actividad.nombreCorto),
            ('Descripción', self.actividad.descripcion),
            ('Tipo de Actividad', self.get_tipo_display()),  # ✅ LLAMADA CORREGIDA
            ('Estado', self.actividad.get_estado_display() if self.actividad.estado else 'N/A'),
            ('Objetivo', self.actividad.objetivo_de_actividad),
            ('Supuestos', self.actividad.supuestos),
            ('Riesgos', self.actividad.riesgos),
        ]
        
        for label, value in datos_basicos:
            if value:
                row = table.add_row()
                row.cells[0].text = str(label)
                row.cells[1].text = str(value)
        
        self.document.add_paragraph()
    
    def agregar_informacion_financiera(self):
        """Agrega información financiera"""
        self.document.add_heading('INFORMACIÓN FINANCIERA', level=1)
        
        table = self.document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        datos_financieros = [
            ('Presupuesto', self.format_currency(self.actividad.presupuesto)),
            ('Presupuesto Global', self.format_currency(self.actividad.presupuestoGlobal)),
            ('Total Reportado', self.format_currency(self.actividad.totalReportado)),
            ('Total Ejecutado', self.format_currency(self.actividad.totalEjecutado)),
            ('Saldo', self.format_currency(self.actividad.saldo)),
            ('Grado de Ejecución', self.actividad.gradoEjecucion),
        ]
        
        for label, value in datos_financieros:
            if value:
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = value
        
        self.document.add_paragraph()
    
    def agregar_fechas_estado(self):
        """Agrega información de fechas y estado"""
        self.document.add_heading('FECHAS Y ESTADO', level=1)
        
        table = self.document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        fechas = [
            ('Fecha Programada', self.format_date(self.actividad.fecha_programada)),
            ('Fecha de Inicio', self.format_date(self.actividad.fecha_inicio)),
            ('Fecha de Cierre', self.format_date(self.actividad.fecha_cierre)),
            ('Estado Actual', self.actividad.get_estado_display() if self.actividad.estado else 'N/A'),
            ('Actividad Inactiva', 'Sí' if self.actividad.estaInactiva else 'No'),
        ]
        
        for label, value in fechas:
            if value:
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = value
        
        self.document.add_paragraph()
    
    def agregar_tareas(self):
        """Agrega información de tareas asociadas"""
        if hasattr(self.actividad, 'tareas') and self.actividad.tareas.exists():
            self.document.add_heading('TAREAS ASOCIADAS', level=1)
            
            for tarea in self.actividad.tareas.all():
                self.document.add_heading(f'Tarea: {tarea.titulo or "Sin título"}', level=2)
                
                table = self.document.add_table(rows=0, cols=2)
                table.style = 'Table Grid'
                
                tarea_datos = [
                    ('Estado', tarea.get_estado_display() if tarea.estado else 'N/A'),
                    ('Descripción', tarea.descripcion),
                    ('Fecha Límite', self.format_date(tarea.fecha_limite)),
                    ('Presupuesto', self.format_currency(tarea.presupuesto)),
                ]
                
                for label, value in tarea_datos:
                    if value:
                        row = table.add_row()
                        row.cells[0].text = label
                        row.cells[1].text = value
                
                self.document.add_paragraph()
    
    def agregar_relaciones(self):
        """Agrega información de relaciones"""
        self.document.add_heading('RELACIONES', level=1)
        
        relaciones = []
        
        if self.actividad.proyecto:
            relaciones.append(('Proyecto', 
                f"{self.actividad.proyecto.codigo or 'N/A'} - {self.actividad.proyecto.titulo or 'N/A'}"))
        
        if self.actividad.responsable:
            # Manejo seguro de atributos del usuario
            username = getattr(self.actividad.responsable, 'username', 'Desconocido')
            
            # Intentar obtener nombre completo
            if hasattr(self.actividad.responsable, 'get_full_name'):
                nombre = self.actividad.responsable.get_full_name()
            else:
                nombre = username
            
            # Intentar obtener email
            email = getattr(self.actividad.responsable, 'email', 'No tiene email')
            
            relaciones.append(('Responsable', f"{nombre} ({email})"))
        
        if self.actividad.proceso:
            relaciones.append(('Proceso', 
                f"{self.actividad.proceso.codigo or 'N/A'} - {self.actividad.proceso.titulo or 'N/A'}"))
        
        if relaciones:
            table = self.document.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            
            for tipo, descripcion in relaciones:
                row = table.add_row()
                row.cells[0].text = tipo
                row.cells[1].text = descripcion
        
        self.document.add_paragraph()
    
    # ✅ MÉTODO FALTANTE AÑADIDO
    def get_tipo_display(self):
        """Obtiene el display del tipo de actividad"""
        if self.actividad.tipo:
            # Verificar si el tipo tiene sigla y tipo_actividad
            sigla = getattr(self.actividad.tipo, 'sigla', '')
            tipo_actividad = getattr(self.actividad.tipo, 'tipo_actividad', '')
            
            if sigla and tipo_actividad:
                return f"{sigla} - {tipo_actividad}"
            elif sigla:
                return sigla
            elif tipo_actividad:
                return tipo_actividad
            else:
                return str(self.actividad.tipo)
        return "No definido"
    
    def guardar_documento_temp(self):
        """Guarda el documento en un archivo temporal"""
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
        self.document.save(temp_file.name)
        temp_file.close()
        return temp_file
    
    def format_currency(self, value):
        """Formatea valores monetarios"""
        if value is None:
            return "No definido"
        try:
            return f"${float(value):,.2f}"
        except (ValueError, TypeError):
            return "Formato inválido"
    
    def format_date(self, date):
        """Formatea fechas"""
        if not date:
            return "No definida"
        try:
            return date.strftime("%d/%m/%Y")
        except AttributeError:
            return "Fecha inválida"