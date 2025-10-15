from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import Proyecto

class ProyectoGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_proyecto(self, proyecto_id, incluir_relaciones=True):
        """Genera reporte completo del proyecto (Nivel 1)"""
        try:
            # proyecto = Proyecto.objects.select_related(
            #     'pei', 'programa'
            # ).prefetch_related(
            #     'instancia_gestora', 'procedencia_fondos'
            # ).get(id=proyecto_id)
            proyecto = Proyecto.objects.get(id=proyecto_id)
            
            # Encadenamiento de secciones
            self._agregar_portada(proyecto)
            self._agregar_informacion_basica(proyecto)
            self._agregar_presupuesto_estado(proyecto)
            self._agregar_relaciones(proyecto)
            
            # Punto de extensión para encadenamiento automático
            if incluir_relaciones:
                self._encadenar_niveles_superiores(proyecto)
            
            return self._guardar_documento()
            
        except Proyecto.DoesNotExist:
            raise ValueError(f"Proyecto con ID {proyecto_id} no encontrado")
    
    def _agregar_portada(self, proyecto):
        """Portada del reporte de proyecto"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE PROYECTO', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Proyecto:', proyecto.codigo),
            ('Título:', proyecto.titulo),
            ('Fecha de Creación:', proyecto.fecha_creacion.strftime('%d/%m/%Y') if proyecto.fecha_creacion else ''),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _agregar_informacion_basica(self, proyecto):
        """Sección de información básica del proyecto"""
        self._agregar_seccion('INFORMACIÓN BÁSICA', 1)
        
        datos_basicos = [
            ('Código', proyecto.codigo),
            ('Título', proyecto.titulo),
            ('Descripción', proyecto.descripcion or 'No disponible'),
            ('Fecha de Inicio', proyecto.fecha_inicio.strftime('%d/%m/%Y') if proyecto.fecha_inicio else 'No definida'),
            ('Fecha de Finalización', proyecto.fecha_finalizacion.strftime('%d/%m/%Y') if proyecto.fecha_finalizacion else 'No definida'),
        ]
        
        self._agregar_tabla_datos('Datos Principales', datos_basicos)
    
    def _agregar_presupuesto_estado(self, proyecto):
        """Sección de presupuesto y estado"""
        self._agregar_seccion('PRESUPUESTO Y ESTADO', 1)
        
        # Formatear presupuesto
        presupuesto_formateado = f"${proyecto.presupuesto:,.2f}" if proyecto.presupuesto else 'No definido'
        
        datos_estado = [
            ('Presupuesto', presupuesto_formateado),
            ('Estado', proyecto.get_estado_display()),
            ('Creado por', proyecto.creado_por or 'No especificado'),
        ]
        
        self._agregar_tabla_datos('Estado del Proyecto', datos_estado)
    
    def _agregar_relaciones(self, proyecto):
        """Sección de relaciones del proyecto"""
        self._agregar_seccion('RELACIONES INSTITUCIONALES', 1)
        
        # Instancias Gestoras
        if proyecto.instancia_gestora.exists():
            instancias = [f"{ig.codigo} - {ig.instancia}" for ig in proyecto.instancia_gestora.all()]
            datos_instancias = [('Instancias Gestoras', '\n'.join(instancias))]
            self._agregar_tabla_datos('Instancias Responsables', datos_instancias)
        
        # Procedencia de Fondos
        if proyecto.procedencia_fondos.exists():
            fondos = [f"{pf.sigla} - {pf.financiera}" for pf in proyecto.procedencia_fondos.all()]
            datos_fondos = [('Fuentes de Financiamiento', '\n'.join(fondos))]
            self._agregar_tabla_datos('Financiamiento', datos_fondos)
        
        # PEI y Programa
        relaciones_externas = []
        if proyecto.pei:
            relaciones_externas.append(('PEI Asociado', str(proyecto.pei)))
        if proyecto.programa:
            relaciones_externas.append(('Programa', str(proyecto.programa)))
        
        if relaciones_externas:
            self._agregar_tabla_datos('Relaciones Estratégicas', relaciones_externas)
    
    def _encadenar_niveles_superiores(self, proyecto):
        """Punto de extensión para encadenamiento automático con niveles superiores"""
        # Este método será implementado cuando agreguemos el ChainComposer
        # Por ahora marca el lugar donde se encadenarán los reportes de niveles 2+
        pass
    
    def generar_y_descargar(self, proyecto_id):
        """Método conveniente para generación y descarga directa"""
        buffer = self.generar_reporte_proyecto(proyecto_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_proyecto_{proyecto_id}.docx"'
        
        return response