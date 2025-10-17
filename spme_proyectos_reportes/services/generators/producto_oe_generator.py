from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import ProductoOE

class ProductoOEGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_producto_oe(self, producto_oe_id):
        """Genera reporte individual del producto de objetivo específico"""
        try:
            producto_oe = ProductoOE.objects.select_related(
                'objetivo_especifico',
                'objetivo_especifico__objetivo_general',
                'objetivo_especifico__objetivo_general__proyecto'
            ).prefetch_related(
                'proceso_producto_oe'
            ).get(id=producto_oe_id)
            
            self._agregar_portada(producto_oe)
            self._agregar_informacion_principal(producto_oe)
            self._agregar_procesos(producto_oe)
            self._agregar_analisis_riesgos(producto_oe)
            self._agregar_vinculaciones(producto_oe)
            
            return self._guardar_documento()
            
        except ProductoOE.DoesNotExist:
            raise ValueError(f"Producto OE con ID {producto_oe_id} no encontrado")
    
    def _agregar_portada(self, producto_oe):
        """Portada del reporte"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE PRODUCTO - OBJETIVO ESPECÍFICO', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Producto:', producto_oe.codigo),
            ('Objetivo Específico:', self._obtener_codigo_objetivo_especifico(producto_oe)),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(producto_oe)),
            ('Proyecto:', self._obtener_info_proyecto(producto_oe)),
            ('Estado:', 'Entregado' if producto_oe.entregado else 'Pendiente'),
            ('Número de Procesos:', str(producto_oe.proceso_producto_oe.count())),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_objetivo_especifico(self, producto_oe):
        """Obtiene el código del objetivo específico de forma segura"""
        if producto_oe.objetivo_especifico and producto_oe.objetivo_especifico.codigo:
            return producto_oe.objetivo_especifico.codigo
        return 'No vinculado a objetivo específico'
    
    def _obtener_codigo_objetivo_general(self, producto_oe):
        """Obtiene el código del objetivo general de forma segura"""
        if (producto_oe.objetivo_especifico and 
            producto_oe.objetivo_especifico.objetivo_general and
            producto_oe.objetivo_especifico.objetivo_general.codigo):
            return producto_oe.objetivo_especifico.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, producto_oe):
        """Obtiene información del proyecto de forma segura"""
        if (producto_oe.objetivo_especifico and 
            producto_oe.objetivo_especifico.objetivo_general and
            producto_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = producto_oe.objetivo_especifico.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, producto_oe):
        """Sección de información principal"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', producto_oe.codigo or 'No definido'),
            ('Descripción', producto_oe.descripcion or 'No disponible'),
            ('Tipo', 'Producto de Objetivo Específico'),
            ('Estado', 'Entregado' if producto_oe.entregado else 'Pendiente'),
            ('Número de Procesos', str(producto_oe.proceso_producto_oe.count())),
        ]
        
        self._agregar_tabla_datos('Datos del Producto', datos_principales)
    
    def _agregar_procesos(self, producto_oe):
        """Sección de procesos asociados al producto OE"""
        self._agregar_seccion('PROCESOS ASOCIADOS', 1)
        
        procesos = producto_oe.proceso_producto_oe.all()
        
        if not procesos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido procesos para este producto.').italic = True
            self.document.add_paragraph()
            return
        
        for i, proceso in enumerate(procesos, 1):
            self.document.add_heading(f'Proceso {i}: {proceso.titulo or "Sin título"}', level=2)
            
            datos_proceso = [
                ('Código', proceso.codigo or 'No definido'),
                ('Título', proceso.titulo or 'No disponible'),
                ('Descripción', proceso.descripcion or 'Sin descripción'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Proceso {i}', datos_proceso)
            self.document.add_paragraph()
    
    def _agregar_analisis_riesgos(self, producto_oe):
        """Sección de análisis de riesgos y supuestos"""
        self._agregar_seccion('ANÁLISIS DEL PRODUCTO', 1)
        
        if producto_oe.supuestos:
            self.document.add_heading('Supuestos del Producto', level=2)
            p_supuestos = self.document.add_paragraph(producto_oe.supuestos)
            p_supuestos.style = 'List Bullet'
            self.document.add_paragraph()
        
        if producto_oe.riesgos:
            self.document.add_heading('Riesgos del Producto', level=2)
            p_riesgos = self.document.add_paragraph(producto_oe.riesgos)
            p_riesgos.style = 'List Bullet'
            self.document.add_paragraph()
        
        if not producto_oe.supuestos and not producto_oe.riesgos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido supuestos ni riesgos para este producto.').italic = True
    
    def _agregar_vinculaciones(self, producto_oe):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Objetivo Específico
        if producto_oe.objetivo_especifico:
            desc_oe = producto_oe.objetivo_especifico.descripcion
            vinculaciones.append(('Objetivo Específico', 
                                f"{producto_oe.objetivo_especifico.codigo} - {desc_oe[:100]}..." if desc_oe else producto_oe.objetivo_especifico.codigo))
        else:
            vinculaciones.append(('Objetivo Específico', 'No vinculado'))
        
        # Objetivo General
        if (producto_oe.objetivo_especifico and 
            producto_oe.objetivo_especifico.objetivo_general):
            og = producto_oe.objetivo_especifico.objetivo_general
            desc_og = og.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{og.codigo} - {desc_og[:100]}..." if desc_og else og.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (producto_oe.objetivo_especifico and 
            producto_oe.objetivo_especifico.objetivo_general and
            producto_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = producto_oe.objetivo_especifico.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        # Resumen
        num_procesos = producto_oe.proceso_producto_oe.count()
        vinculaciones.append(('Procesos Asociados', f'{num_procesos} proceso(s)'))
        vinculaciones.append(('Estado', 'Entregado' if producto_oe.entregado else 'Pendiente'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Producto', vinculaciones)
    
    def generar_y_descargar_individual(self, producto_oe_id):
        """Genera y descarga SOLO el producto OE"""
        buffer = self.generar_reporte_producto_oe(producto_oe_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_producto_oe_{producto_oe_id}.docx"'
        
        return response