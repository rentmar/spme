from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import IndicadorObjetivoEspecifico

class IndicadorOEGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_indicador_oe(self, indicador_oe_id):
        """Genera reporte individual del indicador de objetivo específico"""
        try:
            indicador_oe = IndicadorObjetivoEspecifico.objects.select_related(
                'objetivo_especifico',
                'objetivo_especifico__objetivo_general',
                'objetivo_especifico__objetivo_general__proyecto'
            ).get(id=indicador_oe_id)
            
            self._agregar_portada(indicador_oe)
            self._agregar_informacion_principal(indicador_oe)
            self._agregar_metas_cronograma(indicador_oe)
            self._agregar_vinculaciones(indicador_oe)
            
            return self._guardar_documento()
            
        except IndicadorObjetivoEspecifico.DoesNotExist:
            raise ValueError(f"Indicador de Objetivo Específico con ID {indicador_oe_id} no encontrado")
    
    def _agregar_portada(self, indicador_oe):
        """Portada del reporte"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE INDICADOR - OBJETIVO ESPECÍFICO', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Indicador:', indicador_oe.codigo),
            ('Objetivo Específico:', self._obtener_codigo_objetivo_especifico(indicador_oe)),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(indicador_oe)),
            ('Proyecto:', self._obtener_info_proyecto(indicador_oe)),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_objetivo_especifico(self, indicador_oe):
        """Obtiene el código del objetivo específico de forma segura"""
        if indicador_oe.objetivo_especifico and indicador_oe.objetivo_especifico.codigo:
            return indicador_oe.objetivo_especifico.codigo
        return 'No vinculado a objetivo específico'
    
    def _obtener_codigo_objetivo_general(self, indicador_oe):
        """Obtiene el código del objetivo general de forma segura"""
        if (indicador_oe.objetivo_especifico and 
            indicador_oe.objetivo_especifico.objetivo_general and
            indicador_oe.objetivo_especifico.objetivo_general.codigo):
            return indicador_oe.objetivo_especifico.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, indicador_oe):
        """Obtiene información del proyecto de forma segura"""
        if (indicador_oe.objetivo_especifico and 
            indicador_oe.objetivo_especifico.objetivo_general and
            indicador_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = indicador_oe.objetivo_especifico.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, indicador_oe):
        """Sección de información principal"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', indicador_oe.codigo or 'No definido'),
            ('Descripción', indicador_oe.descripcion or 'No disponible'),
            ('Tipo de Redacción', indicador_oe.get_redaccion_display() if indicador_oe.redaccion else 'No especificado'),
            ('Tipo de Indicador', indicador_oe.get_tipo_display() if indicador_oe.tipo else 'No especificado'),
            ('Frecuencia de Medición', indicador_oe.get_frecuencia_display() if indicador_oe.frecuencia else 'No especificada'),
            ('Responsable', indicador_oe.responsable or 'No asignado'),
        ]
        
        self._agregar_tabla_datos('Datos del Indicador', datos_principales)
        
        # Fuente de verificación
        if indicador_oe.fuente_verificacion:
            self.document.add_heading('Fuente de Verificación', level=2)
            p_fuente = self.document.add_paragraph(indicador_oe.fuente_verificacion)
            p_fuente.style = 'List Bullet'
            self.document.add_paragraph()
        
        # Población objetivo
        if indicador_oe.target_poblacion:
            self.document.add_heading('Población Objetivo', level=2)
            p_poblacion = self.document.add_paragraph(indicador_oe.target_poblacion)
            p_poblacion.style = 'List Bullet'
            if indicador_oe.fechaTargetPoblacion:
                p_fecha = self.document.add_paragraph()
                p_fecha.add_run('Fecha de población objetivo: ').bold = True
                p_fecha.add_run(indicador_oe.fechaTargetPoblacion.strftime('%d/%m/%Y'))
            self.document.add_paragraph()
    
    def _agregar_metas_cronograma(self, indicador_oe):
        """Sección de metas y cronograma"""
        self._agregar_seccion('METAS Y CRONOGRAMA', 1)
        
        # Línea base
        datos_linea_base = [
            ('Línea Base', indicador_oe.baseline or 'No definida'),
            ('Fecha Línea Base', indicador_oe.fechaLineaBase.strftime('%d/%m/%Y') if indicador_oe.fechaLineaBase else 'No definida'),
        ]
        
        self._agregar_tabla_datos('Línea Base', datos_linea_base)
        
        # Metas por trimestre
        metas = []
        if indicador_oe.target_q1:
            metas.append(('Meta Q1', indicador_oe.target_q1, indicador_oe.fechaTargetQ1))
        if indicador_oe.target_q2:
            metas.append(('Meta Q2', indicador_oe.target_q2, indicador_oe.fechaTargetQ2))
        if indicador_oe.target_q3:
            metas.append(('Meta Q3', indicador_oe.target_q3, indicador_oe.fechaTargetQ3))
        if indicador_oe.target_q4:
            metas.append(('Meta Q4', indicador_oe.target_q4, indicador_oe.fechaTargetQ4))
        
        if metas:
            self.document.add_heading('Metas Programadas', level=2)
            datos_metas = []
            for periodo, valor, fecha in metas:
                fecha_str = fecha.strftime('%d/%m/%Y') if fecha else 'No definida'
                datos_metas.append((periodo, f"{valor} - {fecha_str}"))
            
            self._agregar_tabla_datos('Metas', datos_metas)
        else:
            p = self.document.add_paragraph()
            p.add_run('No se han definido metas para este indicador.').italic = True
            self.document.add_paragraph()
    
    def _agregar_vinculaciones(self, indicador_oe):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Objetivo Específico
        if indicador_oe.objetivo_especifico:
            desc_oe = indicador_oe.objetivo_especifico.descripcion
            vinculaciones.append(('Objetivo Específico', 
                                f"{indicador_oe.objetivo_especifico.codigo} - {desc_oe[:100]}..." if desc_oe else indicador_oe.objetivo_especifico.codigo))
        else:
            vinculaciones.append(('Objetivo Específico', 'No vinculado'))
        
        # Objetivo General
        if (indicador_oe.objetivo_especifico and 
            indicador_oe.objetivo_especifico.objetivo_general):
            og = indicador_oe.objetivo_especifico.objetivo_general
            desc_og = og.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{og.codigo} - {desc_og[:100]}..." if desc_og else og.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (indicador_oe.objetivo_especifico and 
            indicador_oe.objetivo_especifico.objetivo_general and
            indicador_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = indicador_oe.objetivo_especifico.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Indicador', vinculaciones)
    
    def generar_y_descargar_individual(self, indicador_oe_id):
        """Genera y descarga SOLO el indicador de objetivo específico"""
        buffer = self.generar_reporte_indicador_oe(indicador_oe_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_indicador_oe_{indicador_oe_id}.docx"'
        
        return response