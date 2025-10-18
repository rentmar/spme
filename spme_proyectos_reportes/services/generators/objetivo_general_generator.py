from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import ObjetivoGeneralProyecto

class ObjetivoGeneralGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_objetivo_general(self, objetivo_general_id, incluir_relaciones=True):
        """Genera reporte completo del objetivo general (Nivel 2)"""
        try:
            objetivo_general = ObjetivoGeneralProyecto.objects.select_related(
                'proyecto'
            ).prefetch_related(
                'objetivos_especificos_og',
                'indicador_og',
                'resultados_og'
            ).get(id=objetivo_general_id)
            
            # Encadenamiento de secciones
            self._agregar_portada(objetivo_general)
            self._agregar_informacion_principal(objetivo_general)
            self._agregar_analisis_riesgos(objetivo_general)
            
            # Punto de extensión para encadenamiento automático
            if incluir_relaciones:
                self._encadenar_niveles_inferiores(objetivo_general)
            
            return self._guardar_documento()
            
        except ObjetivoGeneralProyecto.DoesNotExist:
            raise ValueError(f"Objetivo General con ID {objetivo_general_id} no encontrado")
    
    def _agregar_portada(self, objetivo_general):
        """Portada del reporte de objetivo general"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE OBJETIVO GENERAL', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Objetivo:', objetivo_general.codigo),
            ('Proyecto Asociado:', f"{objetivo_general.proyecto.codigo} - {objetivo_general.proyecto.titulo}"),
            ('Descripción:', objetivo_general.descripcion[:100] + "..." if objetivo_general.descripcion and len(objetivo_general.descripcion) > 100 else objetivo_general.descripcion),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _agregar_informacion_principal(self, objetivo_general):
        """Sección de información principal del objetivo general"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', objetivo_general.codigo or 'No definido'),
            ('Descripción Completa', objetivo_general.descripcion or 'No disponible'),
            ('Proyecto Asociado', f"{objetivo_general.proyecto.codigo} - {objetivo_general.proyecto.titulo}"),
        ]
        
        self._agregar_tabla_datos('Datos del Objetivo General', datos_principales)
    
    def _agregar_analisis_riesgos(self, objetivo_general):
        """Sección de análisis de riesgos y supuestos"""
        self._agregar_seccion('ANÁLISIS DE RIESGOS Y SUPUESTOS', 1)
        
        # Supuestos
        if objetivo_general.supuestos:
            p_supuestos = self.document.add_paragraph()
            p_supuestos.add_run('Supuestos: ').bold = True
            p_supuestos.add_run(objetivo_general.supuestos)
            self.document.add_paragraph()
        
        # Riesgos
        if objetivo_general.riesgos:
            p_riesgos = self.document.add_paragraph()
            p_riesgos.add_run('Riesgos Identificados: ').bold = True
            p_riesgos.add_run(objetivo_general.riesgos)
            self.document.add_paragraph()
        
        # Si no hay información de riesgos/supuestos
        if not objetivo_general.supuestos and not objetivo_general.riesgos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido supuestos ni riesgos para este objetivo general.').italic = True
    
    def _encadenar_niveles_inferiores(self, objetivo_general):
        """Punto de extensión para encadenamiento automático con niveles inferiores"""
        # Este método será implementado completamente con el ChainComposer
        
        # Marcadores para los niveles inferiores que se encadenarán
        niveles_inferiores = [
            ('indicadorog', 'Indicadores del Objetivo General'),
            ('objetivoespecificoog', 'Objetivos Específicos Relacionados'),
            ('resultadoog', 'Resultados del Objetivo General')
        ]
        
        for nivel_id, nombre in niveles_inferiores:
            self.document.add_heading(f'SECCIÓN: {nombre}', level=2)
            p = self.document.add_paragraph()
            p.add_run(f'[Aquí se encadenarán los reportes de {nombre}]').italic = True
            self.document.add_paragraph()

    # === MÉTODOS NUEVOS QUE FALTABAN ===
    
    def generar_reporte_individual(self, objetivo_general_id):
        """Genera SOLO el reporte del objetivo general sin encadenar niveles inferiores"""
        return self.generar_reporte_objetivo_general(objetivo_general_id, incluir_relaciones=False)
    
    def generar_y_descargar_individual(self, objetivo_general_id):
        """Genera y descarga SOLO el objetivo general"""
        buffer = self.generar_reporte_individual(objetivo_general_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_objetivo_general_{objetivo_general_id}.docx"'
        
        return response

    def generar_reporte_completo(self, objetivo_general_id):
        """Método alternativo para compatibilidad"""
        return self.generar_reporte_objetivo_general(objetivo_general_id, incluir_relaciones=True)