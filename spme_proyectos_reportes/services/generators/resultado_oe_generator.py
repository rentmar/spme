from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import ResultadoOE

class ResultadoOEGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_resultado_oe(self, resultado_oe_id):
        """Genera reporte individual del resultado de objetivo específico"""
        try:
            resultado_oe = ResultadoOE.objects.select_related(
                'objetivo_especifico',
                'objetivo_especifico__objetivo_general',
                'objetivo_especifico__objetivo_general__proyecto'
            ).prefetch_related(
                'indicador_res_oe',
                'productos_res_oe',
                'proceso_resultado_oe'
            ).get(id=resultado_oe_id)
            
            self._agregar_portada(resultado_oe)
            self._agregar_informacion_principal(resultado_oe)
            self._agregar_indicadores(resultado_oe)
            self._agregar_productos(resultado_oe)
            self._agregar_procesos(resultado_oe)
            self._agregar_analisis_riesgos(resultado_oe)
            self._agregar_vinculaciones(resultado_oe)
            
            return self._guardar_documento()
            
        except ResultadoOE.DoesNotExist:
            raise ValueError(f"Resultado OE con ID {resultado_oe_id} no encontrado")
    
    def _agregar_portada(self, resultado_oe):
        """Portada del reporte"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE RESULTADO - OBJETIVO ESPECÍFICO', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Resultado:', resultado_oe.codigo),
            ('Objetivo Específico:', self._obtener_codigo_objetivo_especifico(resultado_oe)),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(resultado_oe)),
            ('Proyecto:', self._obtener_info_proyecto(resultado_oe)),
            ('Número de Indicadores:', str(resultado_oe.indicador_res_oe.count())),
            ('Número de Productos:', str(resultado_oe.productos_res_oe.count())),
            ('Número de Procesos:', str(resultado_oe.proceso_resultado_oe.count())),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_objetivo_especifico(self, resultado_oe):
        """Obtiene el código del objetivo específico de forma segura"""
        if resultado_oe.objetivo_especifico and resultado_oe.objetivo_especifico.codigo:
            return resultado_oe.objetivo_especifico.codigo
        return 'No vinculado a objetivo específico'
    
    def _obtener_codigo_objetivo_general(self, resultado_oe):
        """Obtiene el código del objetivo general de forma segura"""
        if (resultado_oe.objetivo_especifico and 
            resultado_oe.objetivo_especifico.objetivo_general and
            resultado_oe.objetivo_especifico.objetivo_general.codigo):
            return resultado_oe.objetivo_especifico.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, resultado_oe):
        """Obtiene información del proyecto de forma segura"""
        if (resultado_oe.objetivo_especifico and 
            resultado_oe.objetivo_especifico.objetivo_general and
            resultado_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = resultado_oe.objetivo_especifico.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, resultado_oe):
        """Sección de información principal"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', resultado_oe.codigo or 'No definido'),
            ('Descripción', resultado_oe.descripcion or 'No disponible'),
            ('Tipo', 'Resultado de Objetivo Específico'),
            ('Nivel', 'Resultado Operativo'),
            ('Número de Indicadores', str(resultado_oe.indicador_res_oe.count())),
            ('Número de Productos', str(resultado_oe.productos_res_oe.count())),
            ('Número de Procesos', str(resultado_oe.proceso_resultado_oe.count())),
        ]
        
        self._agregar_tabla_datos('Datos del Resultado', datos_principales)
    
    def _agregar_indicadores(self, resultado_oe):
        """Sección de indicadores asociados al resultado OE"""
        self._agregar_seccion('INDICADORES ASOCIADOS', 1)
        
        indicadores = resultado_oe.indicador_res_oe.all()
        
        if not indicadores:
            p = self.document.add_paragraph()
            p.add_run('No se han definido indicadores para este resultado.').italic = True
            self.document.add_paragraph()
            return
        
        for i, indicador in enumerate(indicadores, 1):
            nombre_indicador = indicador.descripcion[:50] + "..." if indicador.descripcion else f"Indicador {indicador.codigo}"
            self.document.add_heading(f'Indicador {i}: {nombre_indicador}', level=2)
            
            datos_indicador = [
                ('Código', indicador.codigo or 'No definido'),
                ('Descripción', indicador.descripcion or 'Sin descripción'),
                ('Tipo', indicador.get_tipo_display() if indicador.tipo else 'No especificado'),
                ('Frecuencia', indicador.get_frecuencia_display() if indicador.frecuencia else 'No especificada'),
                ('Línea Base', indicador.baseline or 'No definida'),
                ('Meta Q1', indicador.target_q1 or 'No definida'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Indicador {i}', datos_indicador)
            
            if indicador.fuente_verificacion:
                self.document.add_heading(f'Fuentes de Verificación - Indicador {i}', level=3)
                p_fuentes = self.document.add_paragraph(indicador.fuente_verificacion)
                p_fuentes.style = 'List Bullet'
            
            self.document.add_paragraph()
    
    def _agregar_productos(self, resultado_oe):
        """Sección de productos asociados al resultado OE"""
        self._agregar_seccion('PRODUCTOS ASOCIADOS', 1)
        
        productos = resultado_oe.productos_res_oe.all()
        
        if not productos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido productos para este resultado.').italic = True
            self.document.add_paragraph()
            return
        
        for i, producto in enumerate(productos, 1):
            self.document.add_heading(f'Producto {i}: {producto.codigo}', level=2)
            
            datos_producto = [
                ('Código', producto.codigo or 'No definido'),
                ('Descripción', producto.descripcion or 'No disponible'),
                ('Supuestos', producto.supuestos or 'No definidos'),
                ('Riesgos', producto.riesgos or 'No identificados'),
                ('Entregado', 'Sí' if producto.entregado else 'No'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Producto {i}', datos_producto)
            self.document.add_paragraph()
    
    def _agregar_procesos(self, resultado_oe):
        """Sección de procesos asociados al resultado OE"""
        self._agregar_seccion('PROCESOS ASOCIADOS', 1)
        
        procesos = resultado_oe.proceso_resultado_oe.all()
        
        if not procesos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido procesos para este resultado.').italic = True
            self.document.add_paragraph()
            return
        
        for i, proceso in enumerate(procesos, 1):
            self.document.add_heading(f'Proceso {i}: {proceso.titulo or "Sin título"}', level=2)
            
            datos_proceso = [
                ('Código', proceso.codigo or 'No definido'),
                ('Título', proceso.titulo or 'No disponible'),
                ('Descripción', proceso.descripcion or 'Sin descripción'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Proceso {i}', datos_proceso)
            self.document.add_paragraph()
    
    def _agregar_analisis_riesgos(self, resultado_oe):
        """Sección de análisis de riesgos y supuestos"""
        self._agregar_seccion('ANÁLISIS DEL RESULTADO', 1)
        
        if resultado_oe.supuestos:
            self.document.add_heading('Supuestos del Resultado', level=2)
            p_supuestos = self.document.add_paragraph(resultado_oe.supuestos)
            p_supuestos.style = 'List Bullet'
            self.document.add_paragraph()
        
        if resultado_oe.riesgos:
            self.document.add_heading('Riesgos del Resultado', level=2)
            p_riesgos = self.document.add_paragraph(resultado_oe.riesgos)
            p_riesgos.style = 'List Bullet'
            self.document.add_paragraph()
        
        if not resultado_oe.supuestos and not resultado_oe.riesgos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido supuestos ni riesgos para este resultado.').italic = True
    
    def _agregar_vinculaciones(self, resultado_oe):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Objetivo Específico
        if resultado_oe.objetivo_especifico:
            desc_oe = resultado_oe.objetivo_especifico.descripcion
            vinculaciones.append(('Objetivo Específico', 
                                f"{resultado_oe.objetivo_especifico.codigo} - {desc_oe[:100]}..." if desc_oe else resultado_oe.objetivo_especifico.codigo))
        else:
            vinculaciones.append(('Objetivo Específico', 'No vinculado'))
        
        # Objetivo General
        if (resultado_oe.objetivo_especifico and 
            resultado_oe.objetivo_especifico.objetivo_general):
            og = resultado_oe.objetivo_especifico.objetivo_general
            desc_og = og.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{og.codigo} - {desc_og[:100]}..." if desc_og else og.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (resultado_oe.objetivo_especifico and 
            resultado_oe.objetivo_especifico.objetivo_general and
            resultado_oe.objetivo_especifico.objetivo_general.proyecto):
            proyecto = resultado_oe.objetivo_especifico.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        # Resumen de componentes
        num_indicadores = resultado_oe.indicador_res_oe.count()
        num_productos = resultado_oe.productos_res_oe.count()
        num_procesos = resultado_oe.proceso_resultado_oe.count()
        
        vinculaciones.append(('Indicadores Asociados', f'{num_indicadores} indicador(es)'))
        vinculaciones.append(('Productos Asociados', f'{num_productos} producto(s)'))
        vinculaciones.append(('Procesos Asociados', f'{num_procesos} proceso(s)'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Resultado', vinculaciones)
    
    def generar_y_descargar_individual(self, resultado_oe_id):
        """Genera y descarga SOLO el resultado OE"""
        buffer = self.generar_reporte_resultado_oe(resultado_oe_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_resultado_oe_{resultado_oe_id}.docx"'
        
        return response