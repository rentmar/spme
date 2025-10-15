from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import IndicadorResultadoObjGral

class IndicadorROGGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_indicador_rog(self, indicador_rog_id):
        """Genera reporte individual del indicador de resultado OG"""
        try:
            indicador_rog = IndicadorResultadoObjGral.objects.select_related(
                'resultado_og',
                'resultado_og__objetivo_general',
                'resultado_og__objetivo_general__proyecto'
            ).get(id=indicador_rog_id)
            
            self._agregar_portada(indicador_rog)
            self._agregar_informacion_principal(indicador_rog)
            self._agregar_metas_indicador(indicador_rog)
            self._agregar_vinculaciones(indicador_rog)
            
            return self._guardar_documento()
            
        except IndicadorResultadoObjGral.DoesNotExist:
            raise ValueError(f"Indicador ROG con ID {indicador_rog_id} no encontrado")
    
    def _agregar_portada(self, indicador_rog):
        """Portada del reporte de indicador ROG"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE INDICADOR - RESULTADO OG', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Indicador:', indicador_rog.codigo),
            ('Resultado OG:', self._obtener_codigo_resultado_og(indicador_rog)),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(indicador_rog)),
            ('Proyecto:', self._obtener_info_proyecto(indicador_rog)),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_resultado_og(self, indicador_rog):
        """Obtiene el código del resultado OG de forma segura"""
        if indicador_rog.resultado_og and indicador_rog.resultado_og.codigo:
            return indicador_rog.resultado_og.codigo
        return 'No vinculado a resultado OG'
    
    def _obtener_codigo_objetivo_general(self, indicador_rog):
        """Obtiene el código del objetivo general de forma segura"""
        if (indicador_rog.resultado_og and 
            indicador_rog.resultado_og.objetivo_general and
            indicador_rog.resultado_og.objetivo_general.codigo):
            return indicador_rog.resultado_og.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, indicador_rog):
        """Obtiene información del proyecto de forma segura"""
        if (indicador_rog.resultado_og and 
            indicador_rog.resultado_og.objetivo_general and
            indicador_rog.resultado_og.objetivo_general.proyecto):
            proyecto = indicador_rog.resultado_og.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, indicador_rog):
        """Sección de información principal del indicador"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', indicador_rog.codigo or 'No definido'),
            ('Descripción', indicador_rog.descripcion or 'No disponible'),
            ('Tipo de Indicador', indicador_rog.get_redaccion_display() if indicador_rog.redaccion else 'No definido'),
            ('Tipo de Medición', indicador_rog.get_tipo_display() if indicador_rog.tipo else 'No definido'),
            ('Frecuencia', indicador_rog.get_frecuencia_display() if indicador_rog.frecuencia else 'No definida'),
            ('Fuente de Verificación', indicador_rog.fuente_verificacion or 'No definida'),
            ('Responsable', indicador_rog.responsable or 'No asignado')
        ]
        
        self._agregar_tabla_datos('Datos del Indicador', datos_principales)
    
    def _agregar_metas_indicador(self, indicador_rog):
        """Sección de metas y líneas base del indicador"""
        self._agregar_seccion('METAS Y LÍNEAS BASE', 1)
        
        # Línea Base
        if indicador_rog.baseline or indicador_rog.fechaLineaBase:
            self.document.add_heading('Línea Base', level=2)
            
            datos_linea_base = [
                ('Valor', indicador_rog.baseline or 'No definido'),
                ('Fecha', indicador_rog.fechaLineaBase.strftime('%d/%m/%Y') if indicador_rog.fechaLineaBase else 'No definida')
            ]
            
            tabla_linea_base = self.document.add_table(rows=2, cols=2)
            tabla_linea_base.style = 'Light Grid Accent 1'
            
            for i, (campo, valor) in enumerate(datos_linea_base):
                tabla_linea_base.cell(i, 0).text = campo
                tabla_linea_base.cell(i, 1).text = str(valor)
                tabla_linea_base.cell(i, 0).paragraphs[0].runs[0].bold = True
            
            self.document.add_paragraph()
        
        # Metas por trimestre
        metas_existen = any([
            indicador_rog.target_q1, indicador_rog.target_q2,
            indicador_rog.target_q3, indicador_rog.target_q4
        ])
        
        if metas_existen:
            self.document.add_heading('Metas Programadas', level=2)
            
            tabla_metas = self.document.add_table(rows=5, cols=3)
            tabla_metas.style = 'Medium List 1 Accent 1'
            
            # Encabezados
            headers = ['Período', 'Meta', 'Fecha Meta']
            for i, header in enumerate(headers):
                tabla_metas.cell(0, i).text = header
                tabla_metas.cell(0, i).paragraphs[0].runs[0].bold = True
            
            # Datos de metas
            metas = [
                ('Q1', indicador_rog.target_q1, indicador_rog.fechaTargetQ1),
                ('Q2', indicador_rog.target_q2, indicador_rog.fechaTargetQ2),
                ('Q3', indicador_rog.target_q3, indicador_rog.fechaTargetQ3),
                ('Q4', indicador_rog.target_q4, indicador_rog.fechaTargetQ4)
            ]
            
            for i, (periodo, meta, fecha) in enumerate(metas, 1):
                tabla_metas.cell(i, 0).text = periodo
                tabla_metas.cell(i, 1).text = meta or 'No definida'
                tabla_metas.cell(i, 2).text = fecha.strftime('%d/%m/%Y') if fecha else 'No definida'
            
            self.document.add_paragraph()
        
        # Si no hay metas definidas
        if not indicador_rog.baseline and not metas_existen:
            p = self.document.add_paragraph()
            p.add_run('No se han definido líneas base ni metas para este indicador.').italic = True
    
    def _agregar_vinculaciones(self, indicador_rog):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Resultado OG
        if indicador_rog.resultado_og:
            desc_resultado = indicador_rog.resultado_og.descripcion
            vinculaciones.append(('Resultado OG', 
                                f"{indicador_rog.resultado_og.codigo} - {desc_resultado[:50]}..." if desc_resultado else indicador_rog.resultado_og.codigo))
        else:
            vinculaciones.append(('Resultado OG', 'No vinculado'))
        
        # Objetivo General
        if (indicador_rog.resultado_og and 
            indicador_rog.resultado_og.objetivo_general):
            og = indicador_rog.resultado_og.objetivo_general
            desc_og = og.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{og.codigo} - {desc_og[:50]}..." if desc_og else og.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (indicador_rog.resultado_og and 
            indicador_rog.resultado_og.objetivo_general and
            indicador_rog.resultado_og.objetivo_general.proyecto):
            proyecto = indicador_rog.resultado_og.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Indicador', vinculaciones)
    
    def generar_y_descargar_individual(self, indicador_rog_id):
        """Genera y descarga SOLO el indicador ROG"""
        buffer = self.generar_reporte_indicador_rog(indicador_rog_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_indicador_rog_{indicador_rog_id}.docx"'
        
        return response