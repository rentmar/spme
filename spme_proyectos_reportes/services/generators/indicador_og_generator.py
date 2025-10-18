from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import IndicadorObjetivoGeneral

class IndicadorOGGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_indicador_og(self, indicador_og_id):
        """Genera reporte individual del indicador de objetivo general"""
        try:
            indicador_og = IndicadorObjetivoGeneral.objects.select_related(
                'objetivo_general',
                'objetivo_general__proyecto'
            ).get(id=indicador_og_id)
            
            self._agregar_portada(indicador_og)
            self._agregar_informacion_principal(indicador_og)
            self._agregar_metas_indicador(indicador_og)
            self._agregar_vinculaciones(indicador_og)
            
            return self._guardar_documento()
            
        except IndicadorObjetivoGeneral.DoesNotExist:
            raise ValueError(f"Indicador OG con ID {indicador_og_id} no encontrado")
    
    def _agregar_portada(self, indicador_og):
        """Portada del reporte de indicador OG"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE INDICADOR - OBJETIVO GENERAL', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Indicador:', indicador_og.codigo),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(indicador_og)),
            ('Proyecto:', self._obtener_info_proyecto(indicador_og)),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_objetivo_general(self, indicador_og):
        """Obtiene el código del objetivo general de forma segura"""
        if indicador_og.objetivo_general and indicador_og.objetivo_general.codigo:
            return indicador_og.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, indicador_og):
        """Obtiene información del proyecto de forma segura"""
        if (indicador_og.objetivo_general and 
            indicador_og.objetivo_general.proyecto):
            proyecto = indicador_og.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, indicador_og):
        """Sección de información principal del indicador"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', indicador_og.codigo or 'No definido'),
            ('Descripción', indicador_og.descripcion or 'No disponible'),
            ('Tipo de Indicador', indicador_og.get_redaccion_display() if indicador_og.redaccion else 'No definido'),
            ('Tipo de Medición', indicador_og.get_tipo_display() if indicador_og.tipo else 'No definido'),
            ('Frecuencia', indicador_og.get_frecuencia_display() if indicador_og.frecuencia else 'No definida'),
            ('Fuente de Verificación', indicador_og.fuente_verificacion or 'No definida'),
            ('Responsable', indicador_og.responsable or 'No asignado'),
            ('Población Objetivo', indicador_og.target_poblacion or 'No definida')
        ]
        
        self._agregar_tabla_datos('Datos del Indicador', datos_principales)
    
    def _agregar_metas_indicador(self, indicador_og):
        """Sección de metas y líneas base del indicador"""
        self._agregar_seccion('METAS Y LÍNEAS BASE', 1)
        
        # Línea Base
        if indicador_og.baseline or indicador_og.fechaLineaBase:
            self.document.add_heading('Línea Base', level=2)
            
            datos_linea_base = [
                ('Valor', indicador_og.baseline or 'No definido'),
                ('Fecha', indicador_og.fechaLineaBase.strftime('%d/%m/%Y') if indicador_og.fechaLineaBase else 'No definida')
            ]
            
            tabla_linea_base = self.document.add_table(rows=2, cols=2)
            tabla_linea_base.style = 'Light Grid Accent 1'
            
            for i, (campo, valor) in enumerate(datos_linea_base):
                tabla_linea_base.cell(i, 0).text = campo
                tabla_linea_base.cell(i, 1).text = str(valor)
                tabla_linea_base.cell(i, 0).paragraphs[0].runs[0].bold = True
            
            self.document.add_paragraph()
        
        # Metas por trimestre
        metas_existen = any([
            indicador_og.target_q1, indicador_og.target_q2,
            indicador_og.target_q3, indicador_og.target_q4
        ])
        
        if metas_existen:
            self.document.add_heading('Metas Programadas', level=2)
            
            tabla_metas = self.document.add_table(rows=5, cols=3)
            tabla_metas.style = 'Medium List 1 Accent 1'
            
            # Encabezados
            headers = ['Período', 'Meta', 'Fecha Meta']
            for i, header in enumerate(headers):
                tabla_metas.cell(0, i).text = header
                tabla_metas.cell(0, i).paragraphs[0].runs[0].bold = True
            
            # Datos de metas
            metas = [
                ('Q1', indicador_og.target_q1, indicador_og.fechaTargetQ1),
                ('Q2', indicador_og.target_q2, indicador_og.fechaTargetQ2),
                ('Q3', indicador_og.target_q3, indicador_og.fechaTargetQ3),
                ('Q4', indicador_og.target_q4, indicador_og.fechaTargetQ4)
            ]
            
            for i, (periodo, meta, fecha) in enumerate(metas, 1):
                tabla_metas.cell(i, 0).text = periodo
                tabla_metas.cell(i, 1).text = meta or 'No definida'
                tabla_metas.cell(i, 2).text = fecha.strftime('%d/%m/%Y') if fecha else 'No definida'
            
            self.document.add_paragraph()
        
        # Si no hay metas definidas
        if not indicador_og.baseline and not metas_existen:
            p = self.document.add_paragraph()
            p.add_run('No se han definido líneas base ni metas para este indicador.').italic = True
    
    def _agregar_vinculaciones(self, indicador_og):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Objetivo General
        if indicador_og.objetivo_general:
            desc_og = indicador_og.objetivo_general.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{indicador_og.objetivo_general.codigo} - {desc_og[:50]}..." if desc_og else indicador_og.objetivo_general.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (indicador_og.objetivo_general and 
            indicador_og.objetivo_general.proyecto):
            proyecto = indicador_og.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Indicador', vinculaciones)
    
    def generar_y_descargar_individual(self, indicador_og_id):
        """Genera y descarga SOLO el indicador OG"""
        buffer = self.generar_reporte_indicador_og(indicador_og_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_indicador_og_{indicador_og_id}.docx"'
        
        return response