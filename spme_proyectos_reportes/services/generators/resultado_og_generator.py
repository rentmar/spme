from django.http import HttpResponse
from ..base.base_report_service import BaseReportService
from spme_estructuracion_proyecto.models import (
    ResultadoOG, 
    IndicadorResultadoObjGral,
    Proceso
)

class ResultadoOGGenerator(BaseReportService):
    def __init__(self):
        super().__init__()
    
    def generar_reporte_resultado_og(self, resultado_og_id):
        """Genera reporte individual del resultado de objetivo general"""
        try:
            resultado_og = ResultadoOG.objects.select_related(
                'objetivo_general',
                'objetivo_general__proyecto'
            ).prefetch_related(
                'indicador_res_og',  # Relación con indicadores de resultado OG
                'proceso_resultado_og'  # Relación con procesos
            ).get(id=resultado_og_id)
            
            self._agregar_portada(resultado_og)
            self._agregar_informacion_principal(resultado_og)
            self._agregar_indicadores(resultado_og)
            self._agregar_procesos(resultado_og)
            self._agregar_analisis_riesgos(resultado_og)
            self._agregar_vinculaciones(resultado_og)
            
            return self._guardar_documento()
            
        except ResultadoOG.DoesNotExist:
            raise ValueError(f"Resultado OG con ID {resultado_og_id} no encontrado")
    
    def _agregar_indicadores(self, resultado_og):
        """Sección de indicadores asociados al resultado OG"""
        self._agregar_seccion('INDICADORES ASOCIADOS', 1)
        
        # Obtener indicadores relacionados usando la relación correcta
        indicadores = resultado_og.indicador_res_og.all()
        
        if not indicadores:
            p = self.document.add_paragraph()
            p.add_run('No se han definido indicadores para este resultado.').italic = True
            self.document.add_paragraph()
            return
        
        for i, indicador in enumerate(indicadores, 1):
            # Título del indicador
            nombre_indicador = indicador.descripcion[:50] + "..." if indicador.descripcion else f"Indicador {indicador.codigo}"
            titulo_indicador = self.document.add_heading(f'Indicador {i}: {nombre_indicador}', level=2)
            
            # Datos básicos del indicador - CORREGIDOS según el modelo real
            datos_indicador = [
                ('Código', indicador.codigo or 'No definido'),
                ('Descripción', indicador.descripcion or 'Sin descripción'),
                ('Tipo de Redacción', indicador.get_redaccion_display() if indicador.redaccion else 'No especificado'),
                ('Tipo', indicador.get_tipo_display() if indicador.tipo else 'No especificado'),
                ('Frecuencia', indicador.get_frecuencia_display() if indicador.frecuencia else 'No especificada'),
                ('Línea Base', indicador.baseline or 'No definida'),
                ('Fecha Línea Base', indicador.fechaLineaBase.strftime('%d/%m/%Y') if indicador.fechaLineaBase else 'No definida'),
                ('Meta Q1', indicador.target_q1 or 'No definida'),
                ('Fecha Meta Q1', indicador.fechaTargetQ1.strftime('%d/%m/%Y') if indicador.fechaTargetQ1 else 'No definida'),
                ('Meta Q2', indicador.target_q2 or 'No definida'),
                ('Fecha Meta Q2', indicador.fechaTargetQ2.strftime('%d/%m/%Y') if indicador.fechaTargetQ2 else 'No definida'),
                ('Meta Q3', indicador.target_q3 or 'No definida'),
                ('Fecha Meta Q3', indicador.fechaTargetQ3.strftime('%d/%m/%Y') if indicador.fechaTargetQ3 else 'No definida'),
                ('Meta Q4', indicador.target_q4 or 'No definida'),
                ('Fecha Meta Q4', indicador.fechaTargetQ4.strftime('%d/%m/%Y') if indicador.fechaTargetQ4 else 'No definida'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Indicador {i}', datos_indicador)
            
            # Fuentes de verificación
            if indicador.fuente_verificacion:
                self.document.add_heading(f'Fuentes de Verificación - Indicador {i}', level=3)
                p_fuentes = self.document.add_paragraph(indicador.fuente_verificacion)
                p_fuentes.style = 'List Bullet'
            
            # Población objetivo
            if indicador.target_poblacion:
                self.document.add_heading(f'Población Objetivo - Indicador {i}', level=3)
                p_poblacion = self.document.add_paragraph(indicador.target_poblacion)
                p_poblacion.style = 'List Bullet'
                if indicador.fechaTargetPoblacion:
                    p_fecha = self.document.add_paragraph()
                    p_fecha.add_run('Fecha población objetivo: ').bold = True
                    p_fecha.add_run(indicador.fechaTargetPoblacion.strftime('%d/%m/%Y'))
            
            self.document.add_paragraph()  # Espacio entre indicadores
    
    def _agregar_procesos(self, resultado_og):
        """Sección de procesos asociados al resultado OG"""
        self._agregar_seccion('PROCESOS ASOCIADOS', 1)
        
        # Obtener procesos relacionados
        procesos = resultado_og.proceso_resultado_og.all()
        
        if not procesos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido procesos para este resultado.').italic = True
            self.document.add_paragraph()
            return
        
        for i, proceso in enumerate(procesos, 1):
            # Título del proceso
            titulo_proceso = self.document.add_heading(f'Proceso {i}: {proceso.titulo or "Sin título"}', level=2)
            
            # Datos del proceso
            datos_proceso = [
                ('Código', proceso.codigo or 'No definido'),
                ('Título', proceso.titulo or 'No disponible'),
                ('Descripción', proceso.descripcion or 'Sin descripción'),
            ]
            
            self._agregar_tabla_datos(f'Detalles del Proceso {i}', datos_proceso)
            
            self.document.add_paragraph()  # Espacio entre procesos
    
    def _agregar_portada(self, resultado_og):
        """Portada del reporte de resultado OG"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        titulo = self.document.add_heading('REPORTE DE RESULTADO - OBJETIVO GENERAL', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.document.add_paragraph()
        
        # Información principal de portada
        info_portada = [
            ('Código del Resultado:', resultado_og.codigo),
            ('Objetivo General:', self._obtener_codigo_objetivo_general(resultado_og)),
            ('Proyecto:', self._obtener_info_proyecto(resultado_og)),
            ('Número de Indicadores:', str(resultado_og.indicador_res_og.count())),
            ('Número de Procesos:', str(resultado_og.proceso_resultado_og.count())),
        ]
        
        for etiqueta, valor in info_portada:
            if valor:
                p = self.document.add_paragraph()
                p.add_run(etiqueta).bold = True
                p.add_run(f' {valor}')
        
        self.document.add_page_break()
    
    def _obtener_codigo_objetivo_general(self, resultado_og):
        """Obtiene el código del objetivo general de forma segura"""
        if resultado_og.objetivo_general and resultado_og.objetivo_general.codigo:
            return resultado_og.objetivo_general.codigo
        return 'No vinculado a objetivo general'
    
    def _obtener_info_proyecto(self, resultado_og):
        """Obtiene información del proyecto de forma segura"""
        if (resultado_og.objetivo_general and 
            resultado_og.objetivo_general.proyecto):
            proyecto = resultado_og.objetivo_general.proyecto
            return f"{proyecto.codigo} - {proyecto.titulo}"
        return 'Proyecto no asignado'
    
    def _agregar_informacion_principal(self, resultado_og):
        """Sección de información principal"""
        self._agregar_seccion('INFORMACIÓN PRINCIPAL', 1)
        
        datos_principales = [
            ('Código', resultado_og.codigo or 'No definido'),
            ('Descripción', resultado_og.descripcion or 'No disponible'),
            ('Tipo', 'Resultado de Objetivo General'),
            ('Nivel', 'Resultado Estratégico'),
            ('Número de Indicadores', str(resultado_og.indicador_res_og.count())),
            ('Número de Procesos', str(resultado_og.proceso_resultado_og.count())),
        ]
        
        self._agregar_tabla_datos('Datos del Resultado', datos_principales)
    
    def _agregar_analisis_riesgos(self, resultado_og):
        """Sección de análisis de riesgos y supuestos"""
        self._agregar_seccion('ANÁLISIS DEL RESULTADO', 1)
        
        # Supuestos
        if resultado_og.supuestos:
            self.document.add_heading('Supuestos del Resultado', level=2)
            p_supuestos = self.document.add_paragraph(resultado_og.supuestos)
            p_supuestos.style = 'List Bullet'
            self.document.add_paragraph()
        
        # Riesgos
        if resultado_og.riesgos:
            self.document.add_heading('Riesgos del Resultado', level=2)
            p_riesgos = self.document.add_paragraph(resultado_og.riesgos)
            p_riesgos.style = 'List Bullet'
            self.document.add_paragraph()
        
        # Si no hay información
        if not resultado_og.supuestos and not resultado_og.riesgos:
            p = self.document.add_paragraph()
            p.add_run('No se han definido supuestos ni riesgos para este resultado.').italic = True
    
    def _agregar_vinculaciones(self, resultado_og):
        """Sección de vinculaciones"""
        self._agregar_seccion('VINCULACIONES', 1)
        
        vinculaciones = []
        
        # Objetivo General
        if resultado_og.objetivo_general:
            desc_og = resultado_og.objetivo_general.descripcion
            vinculaciones.append(('Objetivo General', 
                                f"{resultado_og.objetivo_general.codigo} - {desc_og[:100]}..." if desc_og else resultado_og.objetivo_general.codigo))
        else:
            vinculaciones.append(('Objetivo General', 'No vinculado'))
        
        # Proyecto
        if (resultado_og.objetivo_general and 
            resultado_og.objetivo_general.proyecto):
            proyecto = resultado_og.objetivo_general.proyecto
            vinculaciones.append(('Proyecto', 
                                f"{proyecto.codigo} - {proyecto.titulo}"))
        else:
            vinculaciones.append(('Proyecto', 'No asignado'))
        
        # Resumen de componentes
        num_indicadores = resultado_og.indicador_res_og.count()
        num_procesos = resultado_og.proceso_resultado_og.count()
        
        vinculaciones.append(('Indicadores Asociados', f'{num_indicadores} indicador(es)'))
        vinculaciones.append(('Procesos Asociados', f'{num_procesos} proceso(s)'))
        
        if vinculaciones:
            self._agregar_tabla_datos('Relaciones del Resultado', vinculaciones)
    
    def generar_y_descargar_individual(self, resultado_og_id):
        """Genera y descarga SOLO el resultado OG"""
        buffer = self.generar_reporte_resultado_og(resultado_og_id)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="reporte_resultado_og_{resultado_og_id}.docx"'
        
        return response