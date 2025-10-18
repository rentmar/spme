from docx import Document
from docx.shared import Pt

class BaseReportService:
    def __init__(self):
        self.doc = Document()

    def add_title(self, text, level=1):
        if text:
            self.doc.add_heading(str(text), level=level)

    def add_paragraph(self, text):
        if text:
            self.doc.add_paragraph(str(text))

    def add_bullet_list(self, items):
        for item in items:
            self.doc.add_paragraph(str(item), style='List Bullet')

    def add_key_value(self, key, value):
        if value:
            self.doc.add_paragraph(f"{key}: {value}")

    def save(self, filename):
        self.doc.save(filename)
        return filename
