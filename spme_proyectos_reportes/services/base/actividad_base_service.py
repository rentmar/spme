# services/base/actividad_base_service.py
from docx import Document
from docx.shared import Inches, Pt
import io
from datetime import datetime

class ActividadBaseReportService:
    """Servicio base específico para reportes de actividad"""
    
    def __init__(self):
        self.document = Document()
        self._configurar_estilos_base()
    
    def _configurar_estilos_base(self):
        """Configura estilos consistentes para reportes de actividad"""
        # Configurar márgenes
        for section in self.document.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        # Estilos para títulos
        self.document.styles['Heading 1'].font.size = Pt(16)
        self.document.styles['Heading 1'].font.bold = True
        self.document.styles['Heading 2'].font.size = Pt(14)
        self.document.styles['Heading 2'].font.bold = True
        self.document.styles['Heading 3'].font.size = Pt(12)
        self.document.styles['Heading 3'].font.bold = True
    
    def agregar_titulo_principal(self, texto):
        """Agrega un título principal al documento"""
        return self.document.add_heading(texto, level=0)
    
    def agregar_seccion(self, titulo, nivel=1):
        """Agrega una sección con formato consistente"""
        return self.document.add_heading(titulo, level=nivel)
    
    def agregar_parrafo(self, texto, estilo='Normal'):
        """Agrega un párrafo al documento"""
        return self.document.add_paragraph(texto, estilo)
    
    def agregar_tabla_clave_valor(self, datos, titulo=None):
        """Agrega una tabla clave-valor estandarizada"""
        if titulo:
            self.agregar_seccion(titulo, nivel=2)
        
        tabla = self.document.add_table(rows=len(datos), cols=2)
        tabla.style = 'Light Grid Accent 1'
        
        for i, (clave, valor) in enumerate(datos):
            # Celda de clave
            celda_clave = tabla.cell(i, 0)
            celda_clave.text = str(clave)
            celda_clave.paragraphs[0].runs[0].bold = True
            
            # Celda de valor
            celda_valor = tabla.cell(i, 1)
            celda_valor.text = str(valor) if valor is not None else "No definido"
        
        self.agregar_parrafo("")  # Espacio después de la tabla
        return tabla
    
    def formatear_moneda(self, valor):
        """Formatea valores monetarios"""
        if valor is None:
            return "No definido"
        try:
            return f"${float(valor):,.2f}"
        except (TypeError, ValueError):
            return "Formato inválido"
    
    def formatear_fecha(self, fecha):
        """Formatea fechas"""
        if not fecha:
            return "No definida"
        return fecha.strftime("%d/%m/%Y")
    
    def obtener_fecha_actual(self):
        """Retorna la fecha actual formateada"""
        return datetime.now().strftime("%d/%m/%Y %H:%M")
    
    def guardar_documento(self):
        """Guarda el documento en buffer para respuesta HTTP"""
        buffer = io.BytesIO()
        self.document.save(buffer)
        buffer.seek(0)
        return buffer