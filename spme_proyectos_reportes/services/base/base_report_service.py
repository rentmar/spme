from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import io

class BaseReportService:
    def __init__(self):
        self.document = Document()
        self._configurar_estilos_base()
    
    def _configurar_estilos_base(self):
        """Configura estilos consistentes para todos los reportes"""
        # Estilos para títulos
        self.document.styles['Heading 1'].font.size = Pt(16)
        self.document.styles['Heading 1'].font.bold = True
        self.document.styles['Heading 2'].font.size = Pt(14)
        self.document.styles['Heading 2'].font.bold = True
        self.document.styles['Heading 3'].font.size = Pt(12)
        self.document.styles['Heading 3'].font.bold = True
    
    def _agregar_tabla_datos(self, titulo, datos):
        """Método helper para agregar tablas de datos estandarizadas"""
        self.document.add_heading(titulo, level=3)
        
        tabla = self.document.add_table(rows=len(datos), cols=2)
        tabla.style = 'Light Grid Accent 1'
        
        for i, (campo, valor) in enumerate(datos):
            tabla.cell(i, 0).text = str(campo)
            tabla.cell(i, 1).text = str(valor or '')
        
        self.document.add_paragraph()
        return tabla
    
    def _agregar_seccion(self, titulo, nivel=2):
        """Agrega una sección con formato consistente"""
        self.document.add_heading(titulo, level=nivel)
        return self.document.add_paragraph()
    
    def _guardar_documento(self):
        """Guarda el documento en buffer para respuesta HTTP"""
        buffer = io.BytesIO()
        self.document.save(buffer)
        buffer.seek(0)
        return buffer
    
    def limpiar_documento(self):
        """Permite reutilizar el servicio para nuevos reportes"""
        self.document = Document()
        self._configurar_estilos_base()