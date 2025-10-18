from docx import Document
from spme_estructuracion_proyecto.models import Proyecto
from spme_proyectos_reportes.services.objetivo_general_report_service import generar_objetivo_general_reporte
from spme_proyectos_reportes.services.objetivo_especifico_report_service import generar_objetivos_especificos_proyecto

def generar_proyecto_reporte_word(proyecto_id: int, incluir_relaciones=True):
    """
    Genera el reporte Word de un Proyecto (nivel 1)
    e incluye niveles secundarios si se requiere.
    """
    proyecto = Proyecto.objects.prefetch_related(
        "instancia_gestora",
        "procedencia_fondos",
        "objetivo_general__objetivos_especificos_og",
        "objetivo_general__indicador_og",
        "objetivo_general__resultados_og__indicador_res_og",
        "objetivos_especificos"
    ).get(pk=proyecto_id)

    doc = Document()
    doc.add_heading(f"Proyecto: {proyecto.titulo}", level=1)
    doc.add_paragraph(f"Código: {proyecto.codigo}")
    doc.add_paragraph(f"Descripción: {proyecto.descripcion or 'Sin descripción'}")
    doc.add_paragraph(f"Estado: {proyecto.get_estado_display()}")
    if proyecto.presupuesto:
        doc.add_paragraph(f"Presupuesto: ${proyecto.presupuesto:,.2f}")
    doc.add_paragraph()

    doc.add_heading("Instancias Gestoras", level=2)
    if proyecto.instancia_gestora.exists():
        for ig in proyecto.instancia_gestora.all():
            doc.add_paragraph(f"• {ig.codigo or '-'} - {ig.instancia}", style="List Bullet")
    else:
        doc.add_paragraph("No hay instancias gestoras registradas.")

    doc.add_heading("Procedencia de Fondos", level=2)
    if proyecto.procedencia_fondos.exists():
        for pf in proyecto.procedencia_fondos.all():
            doc.add_paragraph(f"• {pf.sigla or '-'} - {pf.financiera}", style="List Bullet")
    else:
        doc.add_paragraph("No se especificó procedencia de fondos.")

    # Encadenar subniveles
    if incluir_relaciones:
        doc.add_page_break()
        generar_objetivo_general_reporte(doc, proyecto)
        generar_objetivos_especificos_proyecto(doc, proyecto)

    return doc
