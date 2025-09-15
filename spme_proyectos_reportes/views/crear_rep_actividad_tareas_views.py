from rest_framework import status
from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import io
from spme_actividades.models import Actividad
from ..serializers.crear_rep_actividad_tareas_serializers import ActividadReporteSerializer 
from django.utils import timezone

#from .models import Actividad
#from .serializers import ActividadReporteSerializer

@api_view(['GET'])
def generar_reporte_actividad_word(request, actividad_id):
    try:
        # Obtener la actividad con todas sus relaciones
        actividad = get_object_or_404(Actividad.objects.prefetch_related(
            'tareas',
            'usuario_actividad_solicitud',
            'usuario_actividad_reembolso',
            'usuario_actividad_sol_pago_directo',
            'usuario_actividad_sol_viaje',
            'rendicion_cuentas_actividad',
            'actividad_informe_actividad'
        ), id=actividad_id)
        
        # Crear documento Word
        document = Document()
        
        # Configuración inicial
        style = document.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(10)
        
        # Título principal
        title = document.add_heading(f'REPORTE COMPLETO DE ACTIVIDAD', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(f'Código: {actividad.codigo} - Nombre: {actividad.nombreCorto}')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Información básica de la actividad
        document.add_heading('1. INFORMACIÓN BÁSICA DE LA ACTIVIDAD', level=1)
        
        # Tabla de información básica
        table_info = document.add_table(rows=1, cols=2)
        table_info.style = 'Table Grid'
        table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Datos básicos
        datos_basicos = [
            ('Código', actividad.codigo),
            ('Nombre', actividad.nombreCorto),
            ('Descripción', actividad.descripcion),
            ('Estado', actividad.estado),
            ('Tipo', actividad.tipo.tipo_actividad if actividad.tipo else 'No definido'),
            ('Fecha Inicio', actividad.fecha_inicio),
            ('Fecha Cierre', actividad.fecha_cierre),
            ('Presupuesto', f"${actividad.presupuesto}" if actividad.presupuesto else 'N/A'),
            ('Presupuesto Global', f"${actividad.presupuestoGlobal}" if actividad.presupuestoGlobal else 'N/A'),
            ('Total Ejecutado', f"${actividad.totalEjecutado}" if actividad.totalEjecutado else 'N/A'),
            ('Saldo', f"${actividad.saldo}" if actividad.saldo else 'N/A'),
            ('Grado Ejecución', actividad.gradoEjecucion),
            ('Responsable', actividad.responsable.username if actividad.responsable else 'No asignado'),
            ('Supuestos', actividad.supuestos),
            ('Riesgos', actividad.riesgos),
            ('Objetivo de Actividad', actividad.objetivo_de_actividad),
        ]
        
        for dato, valor in datos_basicos:
            row = table_info.add_row()
            row.cells[0].text = str(dato)
            row.cells[1].text = str(valor) if valor else 'N/A'

        # Tareas de la actividad - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('2. TAREAS DE LA ACTIVIDAD', level=1)
        if actividad.tareas.exists():
            for i, tarea in enumerate(actividad.tareas.all(), 1):
                document.add_heading(f'2.{i} Tarea: {tarea.titulo}', level=2)
                document.add_paragraph(f'Descripción: {tarea.descripcion}')
                document.add_paragraph(f'Estado: {tarea.get_estado_display()}')
                document.add_paragraph(f'Fecha Límite: {tarea.fecha_limite}')
                document.add_paragraph(f'Presupuesto: ${tarea.presupuesto}' if tarea.presupuesto else 'Presupuesto: N/A')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay tareas registradas para esta actividad.', style='Intense Quote')

        # Solicitudes de Fondos - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('3. SOLICITUDES DE FONDOS', level=1)
        if actividad.usuario_actividad_solicitud.exists():
            for i, solicitud in enumerate(actividad.usuario_actividad_solicitud.all(), 1):
                document.add_heading(f'3.{i} Solicitud: {solicitud.numeroFormulario}', level=2)
                document.add_paragraph(f'Monto Solicitado: ${solicitud.montoSolicitado}' if solicitud.montoSolicitado else 'Monto Solicitado: N/A')
                document.add_paragraph(f'Fecha Solicitud: {solicitud.fechaSolicitud}')
                document.add_paragraph(f'Validación Responsable: {"✅ Aprobado" if solicitud.validacionResponsable else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Coordinador: {"✅ Aprobado" if solicitud.validacionCoordinador else "❌ Pendiente"}')
                document.add_paragraph(f'Lugar: {solicitud.lugarSolicitud}')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay solicitudes de fondos registradas para esta actividad.', style='Intense Quote')

        # Solicitudes de Reembolso - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('4. SOLICITUDES DE REEMBOLSO', level=1)
        if actividad.usuario_actividad_reembolso.exists():
            for i, solicitud in enumerate(actividad.usuario_actividad_reembolso.all(), 1):
                document.add_heading(f'4.{i} Reembolso: {solicitud.numeroFormulario}', level=2)
                document.add_paragraph(f'Monto Solicitado: ${solicitud.montoSolicitado}' if solicitud.montoSolicitado else 'Monto Solicitado: N/A')
                document.add_paragraph(f'Fecha Solicitud: {solicitud.fechaSolicitud}')
                document.add_paragraph(f'Lugar: {solicitud.lugarSolicitud}')
                document.add_paragraph(f'Validación Responsable: {"✅ Aprobado" if solicitud.validacionResponsable else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Coordinador: {"✅ Aprobado" if solicitud.validacionCoordinador else "❌ Pendiente"}')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay solicitudes de reembolso registradas para esta actividad.', style='Intense Quote')

        # Solicitudes de Pago Directo - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('5. SOLICITUDES DE PAGO DIRECTO', level=1)
        if actividad.usuario_actividad_sol_pago_directo.exists():
            for i, solicitud in enumerate(actividad.usuario_actividad_sol_pago_directo.all(), 1):
                document.add_heading(f'5.{i} Pago Directo: {solicitud.numeroFormulario}', level=2)
                document.add_paragraph(f'Nombre: {solicitud.nombre} {solicitud.paterno} {solicitud.materno}')
                document.add_paragraph(f'CI: {solicitud.ci}')
                document.add_paragraph(f'Banco: {solicitud.banco}')
                document.add_paragraph(f'Número de Cuenta: {solicitud.numeroCuenta}')
                document.add_paragraph(f'Monto Solicitado: ${solicitud.montoSolicitado}' if solicitud.montoSolicitado else 'Monto Solicitado: N/A')
                document.add_paragraph(f'Cargo: {solicitud.cargo}')
                document.add_paragraph(f'Validación Responsable: {"✅ Aprobado" if solicitud.validacionResponsable else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Coordinador: {"✅ Aprobado" if solicitud.validacionCoordinador else "❌ Pendiente"}')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay solicitudes de pago directo registradas para esta actividad.', style='Intense Quote')

        # Solicitudes de Viaje - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('6. SOLICITUDES DE VIAJE', level=1)
        if actividad.usuario_actividad_sol_viaje.exists():
            for i, solicitud in enumerate(actividad.usuario_actividad_sol_viaje.all(), 1):
                document.add_heading(f'6.{i} Viaje: {solicitud.numeroFormulario}', level=2)
                document.add_paragraph(f'Evento: {solicitud.evento}')
                document.add_paragraph(f'Fecha Inicio: {solicitud.fechaInicio}')
                document.add_paragraph(f'Fecha Fin: {solicitud.fechaFin}')
                document.add_paragraph(f'Lugar del Evento: {solicitud.lugarEvento}')
                document.add_paragraph(f'Organizador: {solicitud.organizador}')
                document.add_paragraph(f'Monto Solicitado: ${solicitud.montoSolicitado}' if solicitud.montoSolicitado else 'Monto Solicitado: N/A')
                document.add_paragraph(f'Validación Responsable: {"✅ Aprobado" if solicitud.validacionResponsable else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Coordinador: {"✅ Aprobado" if solicitud.validacionCoordinador else "❌ Pendiente"}')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay solicitudes de viaje registradas para esta actividad.', style='Intense Quote')

        # Rendiciones de Cuentas - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('7. RENDICIONES DE CUENTAS', level=1)
        if actividad.rendicion_cuentas_actividad.exists():
            for i, rendicion in enumerate(actividad.rendicion_cuentas_actividad.all(), 1):
                document.add_heading(f'7.{i} Rendición: {rendicion.numeroFormulario}', level=2)
                document.add_paragraph(f'Monto Asignado: ${rendicion.montoAsignado}' if rendicion.montoAsignado else 'Monto Asignado: N/A')
                document.add_paragraph(f'Monto Descargado: ${rendicion.montoDescargado}' if rendicion.montoDescargado else 'Monto Descargado: N/A')
                document.add_paragraph(f'Saldo: ${rendicion.saldo}' if rendicion.saldo else 'Saldo: N/A')
                document.add_paragraph(f'Fecha Desembolso: {rendicion.fechaDesembolso}')
                document.add_paragraph(f'Comprobante Diario: {rendicion.cpteDiario}')
                document.add_paragraph(f'Validación Responsable: {"✅ Aprobado" if rendicion.validacionResponsable else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Coordinador: {"✅ Aprobado" if rendicion.validacionCoordinador else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Contador: {"✅ Aprobado" if rendicion.validacionContador else "❌ Pendiente"}')
                document.add_paragraph(f'Validación Administrador: {"✅ Aprobado" if rendicion.validacionAdministrador else "❌ Pendiente"}')
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay rendiciones de cuentas registradas para esta actividad.', style='Intense Quote')

        # Informes de Actividad - SIEMPRE MOSTRAR TÍTULO
        document.add_heading('8. INFORMES DE ACTIVIDAD', level=1)
        if actividad.actividad_informe_actividad.exists():
            for i, informe in enumerate(actividad.actividad_informe_actividad.all(), 1):
                document.add_heading(f'8.{i} Informe: {informe.numeroInforme}', level=2)
                document.add_paragraph(f'Tipo de Reporte: {informe.reporteTipo}')
                document.add_paragraph(f'Objetivo de la Actividad: {informe.informaObjetivoActividad}')
                document.add_paragraph(f'Herramienta de Evaluación: {informe.herramientaEvaluacion}')
                document.add_paragraph(f'Descripción Medios de Verificación: {informe.descripcionMediosVerificacion}')
                document.add_paragraph(f'Comentarios y Recomendaciones: {informe.comentariosRecomendacion}')
                
                if informe.contribucionesProyecto:
                    document.add_paragraph('Contribuciones al Proyecto:')
                    for contribucion in informe.contribucionesProyecto:
                        document.add_paragraph(f'  • {contribucion}', style='List Bullet')
                
                if informe.contribucionesActividad:
                    document.add_paragraph('Contribuciones a la Actividad:')
                    for contribucion in informe.contribucionesActividad:
                        document.add_paragraph(f'  • {contribucion}', style='List Bullet')
                
                document.add_paragraph('')  # Espacio en blanco
        else:
            document.add_paragraph('No hay informes de actividad registrados para esta actividad.', style='Intense Quote')

        # Pie de página con información de generación
        document.add_paragraph('')
        document.add_paragraph('')
        footer_text = document.add_paragraph('Reporte generado automáticamente el ')
        footer_text.add_run(f'{timezone.now().strftime("%Y-%m-%d %H:%M:%S")}').bold = True
        footer_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Guardar el documento en un buffer
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        
        # Crear respuesta HTTP
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename=reporte_actividad_{actividad.codigo}_{timezone.now().strftime("%Y%m%d_%H%M")}.docx'
        response['Content-Length'] = len(buffer.getvalue())
        
        return response
        
    except Exception as e:
        return Response({
            'error': f'Error al generar el reporte: {str(e)}'
        }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)