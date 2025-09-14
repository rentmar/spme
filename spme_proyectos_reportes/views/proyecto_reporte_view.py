from rest_framework import generics
from rest_framework.response import Response
from rest_framework.decorators import api_view
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import datetime
import io
from spme_estructuracion_proyecto.models import *
from ..serializers.proyecto_reporte_serializer import ProyectoReporteSerializer

class ProyectoReporteCompletoView(generics.RetrieveAPIView):
    """
    Endpoint para generar reporte completo de un proyecto
    """
    queryset = Proyecto.objects.all().select_related(
        'objetivo_general',
        'programa'
    ).prefetch_related(
        'instancia_gestora',
        'procedencia_fondos',
        'objetivo_general__resultados_og',
        'objetivo_general__resultados_og__indicador_res_og',  # Corregido: singular
        'objetivo_general__indicador_og',  # Corregido: singular
        'objetivo_general__objetivos_especificos_og',  # Corregido: relación correcta
        'objetivo_general__objetivos_especificos_og__productos_oe',
        'objetivo_general__objetivos_especificos_og__resultados_oe',
        'objetivo_general__objetivos_especificos_og__resultados_oe__productos_res_oe',
        'objetivo_general__objetivos_especificos_og__resultados_oe__indicador_res_oe',  # Corregido: singular
        'objetivo_general__objetivos_especificos_og__indicador_oe'  # Corregido: singular
    )
    serializer_class = ProyectoReporteSerializer
    lookup_field = 'id'

def generar_reporte_docx(proyecto):
    """Genera un documento DOCX con el reporte del proyecto"""
    document = Document()
    
    # Configuración inicial
    style = document.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Título principal
    title = document.add_heading(f'REPORTE DE PROYECTO: {proyecto.codigo}', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Información básica del proyecto
    document.add_heading('Información Básica del Proyecto', level=1)
    
    info_table = document.add_table(rows=7, cols=2)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ['Código:', proyecto.codigo],
        ['Título:', proyecto.titulo],
        ['Descripción:', proyecto.descripcion or ''],
        ['Estado:', proyecto.get_estado_display()],
        ['Presupuesto:', f"${proyecto.presupuesto or 0:,.2f}"],
        ['Fecha Inicio:', proyecto.fecha_inicio.strftime('%d/%m/%Y') if proyecto.fecha_inicio else ''],
        ['Fecha Fin:', proyecto.fecha_finalizacion.strftime('%d/%m/%Y') if proyecto.fecha_finalizacion else '']
    ]
    
    for i, (label, value) in enumerate(info_data):
        info_table.cell(i, 0).text = label
        info_table.cell(i, 1).text = str(value)
    
    # Objetivo General
    if proyecto.objetivo_general:
        document.add_heading('Objetivo General', level=1)
        og = proyecto.objetivo_general
        document.add_paragraph(f'Código: {og.codigo}')
        document.add_paragraph(f'Descripción: {og.descripcion}')
        
        if og.supuestos:
            document.add_paragraph(f'Supuestos: {og.supuestos}')
        if og.riesgos:
            document.add_paragraph(f'Riesgos: {og.riesgos}')
        
        # Indicadores OG
        if hasattr(og, 'indicador_og') and og.indicador_og.exists():
            document.add_heading('Indicadores del Objetivo General', level=2)
            for indicador in og.indicador_og.all():
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(f'{indicador.codigo}: {indicador.descripcion}')
        
        # Resultados OG
        if og.resultados_og.exists():
            document.add_heading('Resultados del Objetivo General', level=2)
            for resultado in og.resultados_og.all():
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(f'{resultado.codigo}: {resultado.descripcion}')
                
                # Indicadores de Resultados OG
                if hasattr(resultado, 'indicador_res_og') and resultado.indicador_res_og.exists():
                    p = document.add_paragraph()
                    p.add_run('  ○ Indicadores:').bold = True
                    for indicador in resultado.indicador_res_og.all():
                        p = document.add_paragraph()
                        p.add_run('    - ').bold = True
                        p.add_run(f'{indicador.codigo}: {indicador.descripcion}')
        
        # Objetivos Específicos
        if hasattr(og, 'objetivos_especificos_og') and og.objetivos_especificos_og.exists():
            document.add_heading('Objetivos Específicos', level=2)
            for objetivo in og.objetivos_especificos_og.all():
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(f'{objetivo.codigo}: {objetivo.descripcion}')
                
                # Indicadores OE
                if hasattr(objetivo, 'indicador_oe') and objetivo.indicador_oe.exists():
                    p = document.add_paragraph()
                    p.add_run('  ○ Indicadores:').bold = True
                    for indicador in objetivo.indicador_oe.all():
                        p = document.add_paragraph()
                        p.add_run('    - ').bold = True
                        p.add_run(f'{indicador.codigo}: {indicador.descripcion}')
                
                # Productos OE
                if objetivo.productos_oe.exists():
                    p = document.add_paragraph()
                    p.add_run('  ○ Productos:').bold = True
                    for producto in objetivo.productos_oe.all():
                        p = document.add_paragraph()
                        p.add_run('    - ').bold = True
                        p.add_run(f'{producto.codigo}: {producto.descripcion}')
                
                # Resultados OE
                if objetivo.resultados_oe.exists():
                    p = document.add_paragraph()
                    p.add_run('  ○ Resultados:').bold = True
                    for resultado in objetivo.resultados_oe.all():
                        p = document.add_paragraph()
                        p.add_run('    - ').bold = True
                        p.add_run(f'{resultado.codigo}: {resultado.descripcion}')
                        
                        # Indicadores Resultado OE
                        if hasattr(resultado, 'indicador_res_oe') and resultado.indicador_res_oe.exists():
                            p = document.add_paragraph()
                            p.add_run('      ↳ Indicadores:').bold = True
                            for indicador in resultado.indicador_res_oe.all():
                                p = document.add_paragraph()
                                p.add_run('        * ').bold = True
                                p.add_run(f'{indicador.codigo}: {indicador.descripcion}')
                        
                        # Productos Resultado OE
                        if resultado.productos_res_oe.exists():
                            p = document.add_paragraph()
                            p.add_run('      ↳ Productos:').bold = True
                            for producto in resultado.productos_res_oe.all():
                                p = document.add_paragraph()
                                p.add_run('        * ').bold = True
                                p.add_run(f'{producto.codigo}: {producto.descripcion}')
    
    # Información adicional
    document.add_heading('Información Adicional', level=1)
    
    # Instancias Gestoras
    if proyecto.instancia_gestora.exists():
        document.add_heading('Instancias Gestoras', level=2)
        for instancia in proyecto.instancia_gestora.all():
            p = document.add_paragraph()
            p.add_run('• ').bold = True
            p.add_run(f'{instancia.codigo} - {instancia.instancia}')
    
    # Procedencia de Fondos
    if proyecto.procedencia_fondos.exists():
        document.add_heading('Procedencia de Fondos', level=2)
        for fondo in proyecto.procedencia_fondos.all():
            p = document.add_paragraph()
            p.add_run('• ').bold = True
            p.add_run(f'{fondo.sigla} - {fondo.financiera}')
    
    # Programa
    if proyecto.programa:
        document.add_heading('Programa', level=2)
        p = document.add_paragraph()
        p.add_run('• ').bold = True
        p.add_run(f'{proyecto.programa.codigo} - {proyecto.programa.nombre}')
    
    # Fecha de generación
    fecha_gen = document.add_paragraph()
    fecha_gen.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fecha_gen.add_run(f'Reporte generado el: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}').italic = True
    
    return document

@api_view(['GET'])
def descargar_reporte_proyecto(request, id):
    """
    Endpoint para descargar reporte completo en formato DOCX
    """
    proyecto = get_object_or_404(
        Proyecto.objects.select_related(
            'objetivo_general',
            'programa'
        ).prefetch_related(
            'instancia_gestora',
            'procedencia_fondos',
            'objetivo_general__resultados_og',
            'objetivo_general__resultados_og__indicador_res_og',  # Corregido: singular
            'objetivo_general__indicador_og',  # Corregido: singular
            'objetivo_general__objetivos_especificos_og',  # Corregido: relación correcta
            'objetivo_general__objetivos_especificos_og__productos_oe',
            'objetivo_general__objetivos_especificos_og__resultados_oe',
            'objetivo_general__objetivos_especificos_og__resultados_oe__productos_res_oe',
            'objetivo_general__objetivos_especificos_og__resultados_oe__indicador_res_oe',  # Corregido: singular
            'objetivo_general__objetivos_especificos_og__indicador_oe'  # Corregido: singular
        ),
        id=id
    )
    
    # Generar documento DOCX
    doc = generar_reporte_docx(proyecto)
    
    # Guardar en buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Crear respuesta HTTP
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_{proyecto.codigo}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx"'
    
    return response