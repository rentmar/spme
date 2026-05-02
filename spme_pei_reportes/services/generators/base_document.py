#services/generators/base_document.py
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ...utils.word_styles import PAGE_CONFIG

class BaseDocument:
    """Configuración base del documento Word."""
    def __init__(self):
        self.document = Document()
        self._setup_page()
    
    def _setup_page(self):
        section = self.document.sections[0]
        section.top_margin = PAGE_CONFIG['top_margin']
        section.bottom_margin = PAGE_CONFIG['bottom_margin']
        section.left_margin = PAGE_CONFIG['left_margin']
        section.right_margin = PAGE_CONFIG['right_margin']
    
    def agregar_parrafo_estilo(self, texto, estilo_config):
        parrafo = self.document.add_paragraph()
        run = parrafo.add_run(texto)
        if 'font_name' in estilo_config: run.font.name = estilo_config['font_name']
        if 'font_size' in estilo_config: run.font.size = estilo_config['font_size']
        if estilo_config.get('bold'): run.bold = True
        if estilo_config.get('color'): run.font.color.rgb = estilo_config['color']
        if estilo_config.get('italic'): run.italic = True
        if 'alignment' in estilo_config: parrafo.alignment = estilo_config['alignment']
        return parrafo
    
    def crear_tabla_formato(self, filas, columnas, estilo='Light Grid Accent 1'):
        tabla = self.document.add_table(rows=filas, cols=columnas)
        tabla.style = estilo
        return tabla
    
    def agregar_separador(self):
        parrafo = self.document.add_paragraph()
        pPr = parrafo._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '1F4E79')
        pBdr.append(bottom)
        pPr.append(pBdr)
    
    def save_to_bytes(self) -> BytesIO:
        buffer = BytesIO()
        self.document.save(buffer)
        buffer.seek(0)
        return buffer